import sys
sys.stdout.reconfigure(encoding='utf-8')
import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

with open('representants_174_tantan.json', 'r', encoding='utf-8') as f:
    users = json.load(f)

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Représentants Bureaux 174"

# Set headers
headers = ["N°", "Code Bureau", "Commune", "N° Bureau", "Nom d'utilisateur (Login)", "Mot de passe", "Rôle"]
ws.append(headers)

header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
center_align = Alignment(horizontal="center", vertical="center")
left_align = Alignment(horizontal="left", vertical="center")

thin_border = Border(
    left=Side(style='thin', color='D1D5DB'),
    right=Side(style='thin', color='D1D5DB'),
    top=Side(style='thin', color='D1D5DB'),
    bottom=Side(style='thin', color='D1D5DB')
)

for col_idx, h in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col_idx)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = center_align

for idx, u in enumerate(users, 1):
    row = [
        idx,
        u['code_bureau'],
        u['commune'],
        u['numero_bureau'],
        u['username'],
        u['password'],
        'Responsable'
    ]
    ws.append(row)
    current_row = idx + 1
    for col_idx in range(1, 8):
        c = ws.cell(row=current_row, column=col_idx)
        c.border = thin_border
        if col_idx in [1, 2, 4, 7]:
            c.alignment = center_align
        elif col_idx in [5, 6]:
            c.font = Font(name="Consolas", size=11, bold=True)
            c.alignment = center_align
        else:
            c.alignment = left_align

ws.column_dimensions['A'].width = 6
ws.column_dimensions['B'].width = 16
ws.column_dimensions['C'].width = 18
ws.column_dimensions['D'].width = 12
ws.column_dimensions['E'].width = 22
ws.column_dimensions['F'].width = 18
ws.column_dimensions['G'].width = 16

output_path = r'C:\Users\93SAL\Desktop\representants_tan_tan_174.xlsx'
wb.save(output_path)
print(f"✅ Fichier Excel des 174 comptes généré avec succès : {output_path}")
